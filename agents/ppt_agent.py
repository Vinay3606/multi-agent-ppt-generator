from pptx import Presentation
from pathlib import Path


def ppt_agent(file_path):

    file_path = Path(file_path)

    if file_path.suffix.lower() != ".pptx":
        return {
            "status": "error",
            "message": "Only PPTX files are supported"
        }

    presentation = Presentation(file_path)

    slides_data = []

    # Presentation Size
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height

    for index, slide in enumerate(presentation.slides, start=1):

        slide_content = []
        shapes_info = []

        for shape in slide.shapes:

            shape_data = {
                "type": str(shape.shape_type),
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height
            }

            # Text information
            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:

                    slide_content.append(text)

                    shape_data["text"] = text

                    # Font information
                    if shape.has_text_frame:

                        for paragraph in shape.text_frame.paragraphs:

                            for run in paragraph.runs:

                                if run.font.name:
                                    shape_data["font"] = run.font.name

                                if run.font.size:
                                    shape_data["font_size"] = run.font.size.pt

                                break

                            break

            shapes_info.append(shape_data)

        slides_data.append({
            "slide_number": index,
            "content": slide_content,
            "shapes": shapes_info
        })

    return {
        "status": "success",
        "file_name": file_path.name,
        "total_slides": len(presentation.slides),
        "slide_width": slide_width,
        "slide_height": slide_height,
        "slides": slides_data
    }