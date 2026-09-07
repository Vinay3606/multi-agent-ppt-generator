from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pathlib import Path
import re


# =========================================================
# PARSE AI GENERATED CONTENT
# =========================================================

def parse_ppt_content(content):

    slides = []

    # Remove markdown symbols if LLM accidentally generates them
    content = content.replace("**", "")
    content = content.replace("*", "")

    # Find all slides
    slide_pattern = r"(?=^SLIDE\s+\d+\s*$)"

    blocks = re.split(
        slide_pattern,
        content,
        flags=re.MULTILINE | re.IGNORECASE
    )

    for block in blocks:

        block = block.strip()

        if not block:
            continue

        # Get slide number
        slide_match = re.match(
            r"SLIDE\s+(\d+)",
            block,
            re.IGNORECASE
        )

        if not slide_match:
            continue

        slide_number = int(slide_match.group(1))

        slide_data = {
            "slide_number": slide_number,
            "title": "",
            "subtitle": "",
            "description": "",
            "points": []
        }


        # =================================================
        # TITLE
        # =================================================

        title_match = re.search(
            r"^TITLE:\s*(.+)$",
            block,
            re.MULTILINE | re.IGNORECASE
        )

        if title_match:
            slide_data["title"] = title_match.group(1).strip()


        # =================================================
        # SUBTITLE
        # =================================================

        subtitle_match = re.search(
            r"^SUBTITLE:\s*(.+)$",
            block,
            re.MULTILINE | re.IGNORECASE
        )

        if subtitle_match:
            slide_data["subtitle"] = subtitle_match.group(1).strip()


        # =================================================
        # DESCRIPTION
        # =================================================

        description_match = re.search(
            r"^DESCRIPTION:\s*(.*?)(?=^POINT\s+\d+\s+TITLE:|\Z)",
            block,
            re.MULTILINE | re.IGNORECASE | re.DOTALL
        )

        if description_match:

            description = (
                description_match
                .group(1)
                .strip()
                .replace("\n", " ")
            )

            slide_data["description"] = description


        # =================================================
        # POINTS
        # =================================================

        point_pattern = (
            r"^POINT\s+(\d+)\s+TITLE:\s*(.*?)\s*$"
            r"\n"
            r"^POINT\s+\1\s+DESCRIPTION:\s*(.*?)"
            r"(?=^POINT\s+\d+\s+TITLE:|\Z)"
        )

        points = re.findall(
            point_pattern,
            block,
            re.MULTILINE | re.IGNORECASE | re.DOTALL
        )


        for number, point_title, point_description in points:

            slide_data["points"].append({

                "title": (
                    point_title
                    .strip()
                    .replace("\n", " ")
                ),

                "description": (
                    point_description
                    .strip()
                    .replace("\n", " ")
                )

            })


        slides.append(slide_data)


    # Sort slides by number
    slides.sort(
        key=lambda x: x["slide_number"]
    )

    return slides


# =========================================================
# REMOVE ALL OLD TEXT SHAPES
# =========================================================

def remove_old_text(slide):

    shapes_to_remove = []

    for shape in slide.shapes:

        if shape.has_text_frame:

            shapes_to_remove.append(
                shape._element
            )


    for element in shapes_to_remove:

        try:
            element.getparent().remove(element)

        except Exception:
            pass


# =========================================================
# ADD TITLE
# =========================================================

def add_title(slide, title):

    if not title:
        title = "Presentation"


    title_box = slide.shapes.add_textbox(

        Inches(0.8),
        Inches(0.45),
        Inches(11.7),
        Inches(0.8)

    )


    text_frame = title_box.text_frame

    text_frame.word_wrap = True

    paragraph = text_frame.paragraphs[0]

    paragraph.text = title

    paragraph.font.size = Pt(28)

    paragraph.font.bold = True

    paragraph.alignment = PP_ALIGN.LEFT


# =========================================================
# ADD SUBTITLE
# =========================================================

def add_subtitle(slide, subtitle):

    if not subtitle:
        return


    subtitle_box = slide.shapes.add_textbox(

        Inches(0.85),
        Inches(1.2),
        Inches(11.4),
        Inches(0.45)

    )


    text_frame = subtitle_box.text_frame

    text_frame.word_wrap = True

    paragraph = text_frame.paragraphs[0]

    paragraph.text = subtitle

    paragraph.font.size = Pt(15)

    paragraph.alignment = PP_ALIGN.LEFT


# =========================================================
# ADD BODY CONTENT
# =========================================================

def add_body_content(slide, slide_data):


    # =====================================================
    # SLIDE 1 - TITLE SLIDE
    # =====================================================

    if slide_data["slide_number"] == 1:


        if slide_data["description"]:

            body_box = slide.shapes.add_textbox(

                Inches(1.2),
                Inches(2.4),
                Inches(10.9),
                Inches(1.5)

            )


            text_frame = body_box.text_frame

            text_frame.word_wrap = True

            paragraph = text_frame.paragraphs[0]

            paragraph.text = slide_data["description"]

            paragraph.font.size = Pt(20)

            paragraph.alignment = PP_ALIGN.CENTER


        return


    # =====================================================
    # NORMAL SLIDES
    # =====================================================

    body_box = slide.shapes.add_textbox(

        Inches(0.9),
        Inches(1.65),
        Inches(11.4),
        Inches(5.3)

    )


    text_frame = body_box.text_frame

    text_frame.word_wrap = True


    # =====================================================
    # DESCRIPTION
    # =====================================================

    if slide_data["description"]:

        paragraph = text_frame.paragraphs[0]

        paragraph.text = slide_data["description"]

        paragraph.font.size = Pt(17)

        paragraph.space_after = Pt(16)


    # =====================================================
    # POINTS
    # =====================================================

    for point in slide_data["points"]:


        # Point Title
        title_paragraph = text_frame.add_paragraph()

        title_paragraph.text = point["title"]

        title_paragraph.font.size = Pt(17)

        title_paragraph.font.bold = True

        title_paragraph.space_before = Pt(8)

        title_paragraph.space_after = Pt(2)


        # Point Description
        description_paragraph = text_frame.add_paragraph()

        description_paragraph.text = point["description"]

        description_paragraph.font.size = Pt(14)

        description_paragraph.space_after = Pt(8)


# =========================================================
# UPDATE ONE SLIDE
# =========================================================

def update_slide(slide, slide_data):


    # Remove old text only
    remove_old_text(slide)


    # Add title
    add_title(
        slide,
        slide_data["title"]
    )


    # Add subtitle
    add_subtitle(
        slide,
        slide_data["subtitle"]
    )


    # Add body
    add_body_content(
        slide,
        slide_data
    )


# =========================================================
# MAIN PPT GENERATOR
# =========================================================

def generate_ppt(
    template_path,
    generated_content,
    output_path
):


    template_path = Path(template_path)

    output_path = Path(output_path)


    # =====================================================
    # PARSE GENERATED CONTENT
    # =====================================================

    slides_data = parse_ppt_content(
        generated_content
    )


    print("\n" + "=" * 60)

    print(
        f"PARSED {len(slides_data)} SLIDES"
    )

    print("=" * 60)


    for data in slides_data:

        print(
            f"\nSLIDE {data['slide_number']}"
        )

        print(
            "TITLE:",
            data["title"]
        )

        print(
            "SUBTITLE:",
            data["subtitle"]
        )

        print(
            "DESCRIPTION:",
            data["description"]
        )

        print(
            "POINTS:",
            len(data["points"])
        )


    # =====================================================
    # LOAD PPT TEMPLATE
    # =====================================================

    presentation = Presentation(
        template_path
    )


    # =====================================================
    # UPDATE SLIDES
    # =====================================================

    total = min(

        len(presentation.slides),

        len(slides_data)

    )


    for index in range(total):

        print(
            f"\nUpdating Slide {index + 1}"
        )


        slide = presentation.slides[index]

        slide_data = slides_data[index]


        update_slide(

            slide,

            slide_data

        )


    # =====================================================
    # CREATE OUTPUT DIRECTORY
    # =====================================================

    output_path.parent.mkdir(

        parents=True,

        exist_ok=True

    )


    # =====================================================
    # SAVE PPT
    # =====================================================

    presentation.save(
        output_path
    )


    return {

        "status": "success",

        "file": str(output_path),

        "total_slides": len(
            presentation.slides
        )

    }